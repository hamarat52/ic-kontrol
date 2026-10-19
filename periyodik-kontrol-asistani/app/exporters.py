"""Word (.docx) ve Excel (.xlsx) dışa aktarımı. Word: kurumun kontrol kartı şablonu (KN.01 – …)."""
import re
from datetime import datetime
from io import BytesIO

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INK = RGBColor(0x1F, 0x20, 0x25)
MUTED = RGBColor(0x6E, 0x6F, 0x76)
KATEGORI_SIRASI = ["Operasyonel", "Mevzuat Uyum", "Bilgi Sistemleri", "Finansal Raporlama", "AML / Şüpheli İşlem"]
DURUM = {"bekliyor": "Bekliyor", "kabul": "Kabul", "red": "Red"}


def _durum(s):
    st = (s.get("status") or "bekliyor").lower()
    if st.startswith("düzenlendi"):
        return "Düzenlendi"
    return DURUM.get(st, st.capitalize())


def _grouped(controls):
    groups = {}
    for c in controls:
        groups.setdefault(c.get("kategori") or "Diğer", []).append(c)
    order = [k for k in KATEGORI_SIRASI if k in groups] + [k for k in groups if k not in KATEGORI_SIRASI]
    return [(k, groups[k]) for k in order]


def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def build_word(control_set: dict, controls: list[dict], library: bool = False) -> BytesIO:
    """library=True: Kontrol Kütüphanesi çıktısı (Durum yerine kontrolün geldiği set)."""
    doc = DocxDocument()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10.5)
    base.font.color.rgb = INK

    h = doc.add_heading(control_set.get("baslik") or "Periyodik Kontroller", level=0)
    for r in h.runs:
        r.font.color.rgb = INK
        r.font.size = Pt(24)

    meta = doc.add_paragraph()
    belgeler = control_set.get("belge_adlari") or []
    parts = [
        f"{len(belgeler)} belge" if belgeler else None,
        (control_set.get("created_at") or "")[:10] or datetime.now().strftime("%Y-%m-%d"),
        f"{len(controls)} kontrol",
        control_set.get("model"),
    ]
    run = meta.add_run("  ·  ".join(p for p in parts if p))
    run.font.color.rgb = MUTED
    run.font.size = Pt(9.5)

    if control_set.get("ozet"):
        doc.add_paragraph(control_set["ozet"])

    for n, c in enumerate(_ordered(controls), start=1):
        _control_card(doc, c, _code(c, n, library), library)

    gaps = control_set.get("bosluklar") or []
    if gaps:
        doc.add_heading("Tespit edilen boşluklar", level=2)
        for g in gaps:
            doc.add_paragraph(g, style="List Bullet")

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    n = note.add_run("Periyodik Kontrol Asistanı tarafından üretilmiştir; kullanılmadan önce uzman gözden geçirmelidir.")
    n.italic = True
    n.font.size = Pt(8.5)
    n.font.color.rgb = MUTED

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ------------------------------------------------------------------ kurum kontrol kartı (KN.01 – …)

CARD_HEAD = RGBColor(0x1F, 0x38, 0x64)
LABEL_FILL = "DDEDE8"
BORDER = "BFBFBF"
SIKLIK_KISA = {"Altı aylık": "6 aylık", "Üç aylık": "3 aylık"}
COL_W = (Cm(4.3), Cm(12.7))


def _ordered(controls):
    """Kod sırasıyla (KN.01, KN.02…); kodu olmayanlar sonda, kategoriye göre."""
    def key(c):
        m = re.match(r"KN\.(\d+)", c.get("kod") or "")
        return (0, int(m.group(1))) if m else (1, KATEGORI_SIRASI.index(c["kategori"]) if c.get("kategori") in KATEGORI_SIRASI else 9)
    return sorted(controls, key=key)


def _code(c, n, library):
    # Kütüphanede kontroller farklı setlerden gelir: kodlar çakışmasın diye yeniden numaralanır
    return f"KN.{n:02d}" if library or not c.get("kod") else c["kod"]


def _borders(table):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), BORDER)
        b.append(e)
    tbl_pr.append(b)


def _card_rows(c):
    siklik = SIKLIK_KISA.get(c.get("siklik"), c.get("siklik") or "")
    donem = [x.strip() for x in re.split(r"[;\n]+", c.get("donem") or "") if x.strip()]
    adim = [x.strip() for x in (c.get("adimlar") or "").split("\n") if x.strip()]
    return [
        ("Kontrol Amacı", [c.get("amac") or c.get("gerekce") or ""]),
        ("Dayanak", [c.get("dayanak") or "-"]),
        ("Kontrolün Niteliği", [c.get("nitelik") or "-"]),
        ("İlgili Birimler", [c.get("ilgili_birimler") or c.get("sorumlu") or "-"]),
        ("Ana Kütle / Kapsam", [c.get("kapsam") or "-"]),
        ("Sıklık ve Dönem", [f"{siklik}." if siklik else "-"] + [f"• {d}" for d in donem]),
        ("Uygulama Adımları", [f"{i}) {a}" for i, a in enumerate(adim, start=1)] or ["-"]),
        ("Asgari Kanıt", [c.get("kanit") or "-"]),
        ("Değerlendirme Ölçütü", [c.get("olcut") or "-"]),
        ("İstisna / Aksiyon", [c.get("istisna") or "-"]),
        ("Raporlama", [c.get("raporlama") or "-"]),
    ]


def _control_card(doc, c, code, library):
    h = doc.add_heading(f"{code} – {c.get('kontrol_adi') or ''}", level=2)
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.keep_with_next = True
    for r in h.runs:
        r.font.color.rgb = CARD_HEAD
        r.font.size = Pt(12.5)
    rows = _card_rows(c)
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _borders(table)
    for col, w in zip(table.columns, COL_W):     # tblGrid: Word ve LibreOffice sütun genişliğini buradan alır
        col.width = w
    for (label, lines), row in zip(rows, table.rows):
        lc, vc = row.cells
        lc.width, vc.width = COL_W
        _shade(lc, LABEL_FILL)
        lr = lc.paragraphs[0].add_run(label)
        lr.bold = True
        lr.font.size = Pt(9.5)
        first = True
        for line in lines:
            par = vc.paragraphs[0] if first else vc.add_paragraph()
            first = False
            par.paragraph_format.space_after = Pt(1)
            if line.startswith("• "):
                par.paragraph_format.left_indent = Cm(0.3)
            run = par.add_run(line)
            run.font.size = Pt(9.5)
    # Kurum şablonunda olmayan ama uygulamanın ürettiği bilgiler: kartın altında küçük not
    extra = [f"Risk: {c.get('risk')}" if c.get("risk") else None, f"Hat {c.get('hat')}" if c.get("hat") else None,
             f"Sorumlu: {c['sorumlu']}" if c.get("sorumlu") else None]
    if c.get("kaynak_belge"):
        src = f"Kaynak: {c['kaynak_belge']}" + (f" · {c['kaynak_bolum']}" if c.get("kaynak_bolum") else "")
        if c.get("kaynak_dogrulandi") == 0:
            src += " (alıntı belgede bulunamadı)"
        extra.append(src)
    extra.append(f"Set: {c['set_baslik']}" if library and c.get("set_baslik") else (f"Durum: {_durum(c)}" if not library else None))
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(3)
    nr = note.add_run("  ·  ".join(x for x in extra if x))
    nr.font.size = Pt(8)
    nr.font.color.rgb = MUTED


EXCEL_COLS = [
    ("Kod", "kod", 8), ("Kontrol", "kontrol_adi", 48), ("Kontrol amacı", "amac", 50), ("Kategori", "kategori", 20),
    ("Hat", "hat", 6), ("Sorumlu", "sorumlu", 24), ("İlgili birimler", "ilgili_birimler", 30),
    ("Kontrolün niteliği", "nitelik", 36), ("Ana kütle / kapsam", "kapsam", 40), ("Dönem", "donem", 36),
    ("Uygulama adımları", "adimlar", 60), ("Asgari kanıt", "kanit", 40), ("Değerlendirme ölçütü", "olcut", 40),
    ("İstisna / aksiyon", "istisna", 30), ("Raporlama", "raporlama", 30),
    ("Sıklık", "siklik", 16), ("Risk", "risk", 10), ("Efor", "efor", 10),
    ("Gerekçe", "gerekce", 60), ("Dayanak", "dayanak", 30),
    ("Kaynak belge", "kaynak_belge", 28), ("Kaynak bölüm", "kaynak_bolum", 28),
    ("Kaynak alıntı", "kaynak", 50), ("Alıntı doğrulandı", "_dogrulama", 14),
    ("Durum", None, 12),
]


def _dogrulama(c):
    v = c.get("kaynak_dogrulandi")
    return "—" if v is None else ("Evet" if v else "Hayır — belgede bulunamadı")


def build_excel(control_set: dict, controls: list[dict], library: bool = False) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kontroller"
    thin = Side(style="thin", color="E6E5E0")
    head_fill = PatternFill("solid", fgColor="1F2025")

    cols = [(("Kaynak set", "set_baslik", 30) if key is None else (h, key, w)) for h, key, w in EXCEL_COLS] if library else EXCEL_COLS
    ws.append([c[0] for c in cols])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22

    for kategori, items in _grouped(controls):
        for c in items:
            row = []
            for _, key, _ in cols:
                if key is None:
                    row.append(_durum(c))
                elif key == "_dogrulama":
                    row.append(_dogrulama(c))
                elif key == "adimlar":
                    row.append("\n".join(f"{i}) {a}" for i, a in enumerate((c.get("adimlar") or "").split("\n"), 1) if a.strip()))
                else:
                    row.append(c.get(key) or "")
            ws.append(row)

    for i, (_, _, width) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Bilgi" if library else "Set bilgisi")
    rows = [
        ("Başlık", control_set.get("baslik") or ""),
        ("Oluşturulma", control_set.get("created_at") or ""),
        ("Model", control_set.get("model") or ""),
        ("Belgeler", ", ".join(control_set.get("belge_adlari") or [])),
        ("Özet", control_set.get("ozet") or ""),
    ] + [("Boşluk", g) for g in (control_set.get("bosluklar") or [])]
    for r in rows:
        info.append(list(r))
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 100
    for row in info.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
