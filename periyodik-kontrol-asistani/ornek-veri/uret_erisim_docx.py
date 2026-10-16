"""Sentetik 'Kullanıcı Erişim Yönetimi Uygulama Esasları' Word belgesini üretir (yetki tablosu dahil).
Çalıştır: python ornek-veri/uret_erisim_docx.py"""
from pathlib import Path

from docx import Document

doc = Document()
doc.add_heading("Kullanıcı Erişim Yönetimi Uygulama Esasları", level=0)
doc.add_paragraph("ÖRNEK BANKA A.Ş. — Sentetik test belgesidir; gerçek bir kuruma ait değildir.")

sections = {
    "1. Amaç ve Kapsam": [
        "Bu esaslar, bankanın tüm bilgi sistemlerine kullanıcı erişimlerinin tanımlanması, değiştirilmesi ve kaldırılmasına ilişkin kuralları belirler.",
        "Temel bankacılık sistemi, ödeme sistemleri, kredi başvuru sistemi ve doküman yönetim sistemi kapsamdadır.",
    ],
    "2. Kullanıcı Tanımlama": [
        "Yeni kullanıcı talepleri İnsan Kaynakları tarafından işe giriş kaydı oluşturulduktan sonra erişim talep sistemine girilir.",
        "Talep, kullanıcının yöneticisi tarafından onaylanmadan Bilgi Güvenliği birimine iletilmez.",
        "Kullanıcıya yalnız görev tanımının gerektirdiği roller atanır; rol dışı yetki talepleri gerekçeli olarak Bilgi Güvenliği onayına sunulur.",
        "İlk parola kullanıcıya kapalı zarfla veya SMS ile iletilir ve ilk girişte değiştirilmesi zorunludur.",
    ],
    "3. Yetki Matrisi": [
        "Aşağıdaki tablo, kritik işlemler için rol bazlı yetkileri gösterir. Aynı kullanıcıya hem 'giriş' hem 'onay' rolü verilemez.",
    ],
    "4. Periyodik Yetki Gözden Geçirme": [
        "Birim yöneticileri, ekiplerindeki kullanıcıların yetkilerini altı ayda bir gözden geçirir ve onaylar.",
        "Gözden geçirme sonucunda gereksiz bulunan yetkiler on iş günü içinde kaldırılır.",
        "Gözden geçirme kayıtları Bilgi Güvenliği birimi tarafından saklanır.",
    ],
    "5. Görev Değişikliği ve Ayrılan Personel": [
        "Görev değişikliğinde eski göreve ait yetkiler yeni yetkiler tanımlanmadan önce kaldırılır.",
        "İşten ayrılan personelin tüm erişimleri ayrılış tarihinde gün sonuna kadar kapatılır.",
        "İnsan Kaynakları, ayrılan personel listesini her gün Bilgi Güvenliği birimine otomatik olarak iletir.",
    ],
    "6. Ayrıcalıklı Hesaplar": [
        "Veritabanı yöneticisi ve sistem yöneticisi gibi ayrıcalıklı hesaplar kişiye özel tanımlanır; ortak hesap kullanılmaz.",
        "Ayrıcalıklı hesaplarla yapılan tüm işlemler kayıt altına alınır ve Bilgi Güvenliği tarafından haftalık incelenir.",
        "Acil erişim (break-glass) hesabı yalnız kayıtlı bir olay numarasıyla açılır ve kullanım sonrası parolası değiştirilir.",
    ],
    "7. Log İzleme": [
        "Başarısız oturum açma denemeleri beşi aştığında hesap otomatik kilitlenir.",
        "Mesai dışı yapılan kritik işlemler ve yetki değişiklikleri günlük rapor halinde Bilgi Güvenliği birimine sunulur.",
    ],
}

for title, paras in sections.items():
    doc.add_heading(title, level=1)
    for p in paras:
        doc.add_paragraph(p)
    if title.startswith("3."):
        rows = [
            ("İşlem", "Giriş rolü", "Onay rolü", "Limit"),
            ("Kredi tahsisi", "Portföy yöneticisi", "Şube müdürü / Kredi komitesi", "Yetki tablosuna göre"),
            ("EFT / Havale", "Operasyon personeli", "Operasyon yetkilisi", "250.000 TL'ye kadar"),
            ("Parametre değişikliği", "Sistem analisti", "Bilgi Sistemleri müdürü", "—"),
            ("Kullanıcı yetki atama", "Bilgi Güvenliği uzmanı", "Bilgi Güvenliği müdürü", "—"),
        ]
        t = doc.add_table(rows=len(rows), cols=4)
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                t.cell(i, j).text = v

out = Path(__file__).resolve().parent / "sentetik-erisim-yonetimi-uygulama-esaslari.docx"
doc.save(out)
print(out)
