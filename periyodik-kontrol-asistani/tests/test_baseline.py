"""Mevcut davranışın baz testleri (adım 0)."""
import io
import json

from docx import Document
from openpyxl import load_workbook

from app import llm
from conftest import jsonl, wait_set

DOC = ("KREDİ TAHSİS PROSEDÜRÜ\n4.1. Her kredi başvurusu, tahsis onayından önce Risk Analizi Biriminin "
       "değerlendirmesinden geçer.\n5.3. Aynı kullanıcı hem tahsis teklifini hazırlayıp hem onaylayamaz.\n")

SET_OUTPUT = jsonl(
    {"tip": "baslik", "baslik": "Kredi Tahsis", "ozet": "Tahsis süreci incelendi."},
    {"tip": "kontrol", "kontrol_adi": "Risk analizi atlanan onaylar", "kategori": "operasyonel", "hat": "Hat 1",
     "siklik": "günlük", "risk": "yuksek", "efor": "orta", "gerekce": "Onay zinciri", "dayanak": "md. 4.1",
     "kaynak": "tahsis onayından önce Risk Analizi Biriminin değerlendirmesinden geçer"},
    {"tip": "kontrol", "kontrol_adi": "Görevler ayrılığı testi", "kategori": "Bilgi Sistemleri", "hat": "2",
     "siklik": "Aylık", "risk": "Kritik", "efor": "Düşük", "gerekce": "SoD", "dayanak": "", "kaynak": ""},
    {"tip": "bosluk", "metin": "Hızlı onay yolu tanımsız."},
)


def set_responder(text, summary=None):
    summary = summary or json.dumps({"baslik": "Kredi Tahsis", "ozet": "Tahsis süreci incelendi."}, ensure_ascii=False)
    return lambda messages, stream: text if stream else summary


# ------------------------------------------------------------------ birim

def test_pick_normalizes_vocab():
    assert llm._pick("yuksek", llm.RISKLER, None) == "Yüksek"
    assert llm._pick("AML", llm.KATEGORILER, None) == "AML / Şüpheli İşlem"
    assert llm._pick("aylık", llm.SIKLIKLAR, None) == "Aylık"
    assert llm._pick("saçma", llm.EFORLAR, "Orta") == "Orta"


def test_normalize_control_hat_and_defaults():
    c = llm.normalize_control({"kontrol_adi": " X ", "hat": "Hat 3 — iç denetim"})
    assert c["kontrol_adi"] == "X" and c["hat"] == "3" and c["kategori"] == "Operasyonel"


def test_stream_parses_jsonl(env):
    env.responder = set_responder(SET_OUTPUT)
    events = list(llm.stream_control_set(DOC, "skills"))
    assert [e["event"] for e in events] == ["meta", "control", "control", "gap"]
    assert events[1]["control"]["risk"] == "Yüksek"
    assert events[1]["control"]["hat"] == "1"


def test_stream_recovers_from_json_array(env):
    env.responder = set_responder(json.dumps([{"kontrol_adi": "A", "kategori": "Mevzuat"}], ensure_ascii=False))
    events = list(llm.stream_control_set(DOC, "skills"))
    assert [e["event"] for e in events] == ["control"]
    assert events[0]["control"]["kategori"] == "Mevzuat Uyum"


def test_stream_reports_model_failure(env):
    env.status = 500
    events = list(llm.stream_control_set(DOC, "skills"))
    assert events[-1]["event"] == "error"


# ------------------------------------------------------------------ uçtan uca

def _create_set(client, env):
    client.post("/api/documents", files={"file": ("prosedur.txt", DOC.encode(), "text/plain")})
    env.responder = set_responder(SET_OUTPUT)
    r = client.post("/api/sets")
    assert r.status_code == 202 and r.json()["durum"] == "uretiliyor"
    s = wait_set(client, r.json()["id"])
    assert s["durum"] == "hazir"
    return s


def test_full_flow(client, env):
    s = _create_set(client, env)
    set_id = s["id"]
    assert s["baslik"] == "Kredi Tahsis" and s["durum"] == "hazir"
    assert len(s["controls"]) == 2 and s["bosluklar"] == ["Hızlı onay yolu tanımsız."]

    c1 = s["controls"][0]["id"]
    assert client.post(f"/api/controls/{c1}/decision", json={"karar": "kabul"}).json()["control"]["status"] == "kabul"
    assert len(client.get("/api/library").json()) == 1

    r = client.patch(f"/api/controls/{c1}", json={"siklik": "Haftalık"}).json()
    assert r["changed"] and client.get("/api/library").json()[0]["siklik"] == "Haftalık"

    hist = client.get(f"/api/controls/{c1}/history").json()
    assert [h["text"].split(" ·")[0] for h in hist] == ["Model önerdi", "Kabul edildi", "Düzenlendi"]

    client.post(f"/api/controls/{c1}/decision", json={"karar": "bekliyor"})
    assert client.get("/api/library").json() == []


def test_exports_are_valid(client, env):
    set_id = _create_set(client, env)["id"]
    w = client.get(f"/api/export/word?set_id={set_id}")
    assert w.status_code == 200
    doc = Document(io.BytesIO(w.content))
    assert doc.paragraphs[0].text == "Kredi Tahsis" and len(doc.tables) == 2

    x = client.get(f"/api/export/excel?set_id={set_id}")
    wb = load_workbook(io.BytesIO(x.content))
    assert wb["Kontroller"].max_row == 3 and "Set bilgisi" in wb.sheetnames


def test_rejected_controls_not_exported(client, env):
    set_id = _create_set(client, env)["id"]
    ids = [c["id"] for c in client.get(f"/api/sets/{set_id}").json()["controls"]]
    client.post(f"/api/controls/{ids[0]}/decision", json={"karar": "red"})
    wb = load_workbook(io.BytesIO(client.get(f"/api/export/excel?set_id={set_id}").content))
    assert wb["Kontroller"].max_row == 2


def test_chat_applies_changes_and_learns(client, env):
    set_id = _create_set(client, env)["id"]
    cid = client.get(f"/api/sets/{set_id}").json()["controls"][0]["id"]
    env.responder = lambda m, s: json.dumps({
        "yanit": "Sıklığı düşürdüm.",
        "degisiklikler": [{"id": cid, "alanlar": {"siklik": "haftalık", "risk": "uydurma"}},
                          {"id": 99999, "alanlar": {"risk": "Düşük"}}],
        "ders": "Düşük hacimli şubelerde sıklık bir kademe düşürülebilir.",
    }, ensure_ascii=False)
    r = client.post(f"/api/sets/{set_id}/chat", json={"message": "sıklık fazla", "control_id": cid}).json()
    assert r["ok"]
    assert r["message"]["changes"] == [
        {"id": cid, "kontrol_adi": "Risk analizi atlanan onaylar", "alanlar": {"siklik": ["Günlük", "Haftalık"]}}]
    skills = {s["name"]: s["content"] for s in client.get("/api/skills").json()}
    assert "Düşük hacimli şubelerde" in skills["kontrol-metrikleri"]
    assert len(client.get(f"/api/sets/{set_id}").json()["messages"]) == 2


def test_generation_without_documents_is_rejected(client, env):
    assert client.post("/api/sets").status_code == 400


def test_generation_model_down_marks_set_failed(client, env):
    client.post("/api/documents", files={"file": ("p.txt", DOC.encode(), "text/plain")})
    env.status = 503
    s = wait_set(client, client.post("/api/sets").json()["id"])
    assert s["durum"] == "hata" and "işlenemedi" in s["hata"]
