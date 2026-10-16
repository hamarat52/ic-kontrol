"""Adım 1: bölüm bazlı üretim, izlenebilirlik, tekrar ayıklama, iptal, kurtarma."""
import io
import json
import re
import time

from docx import Document

from app import documents, jobs, storage
from conftest import jsonl, wait_set

SECTION = "{n}. Bölüm {n} Başlığı\n" + "\n".join(
    f"{{n}}.{i}. Bölüm {{n}} için kural {i}: işlem yetkili tarafından onaylanır ve kayda alınır." for i in range(1, 40))


def long_doc(sections=4):
    return "\n".join(SECTION.format(n=n) for n in range(1, sections + 1))


def bolum_of(messages):
    user = next(m["content"] for m in messages if m["role"] == "user")
    m = re.search(r"BÖLÜM: (.+?) \((\d+)/(\d+)\)", user)
    return m.group(1), int(m.group(2)), int(m.group(3)), user


def summary(_m=None):
    return json.dumps({"baslik": "Test Seti", "ozet": "Özet."}, ensure_ascii=False)


def upload(client, name, text):
    return client.post("/api/documents", files={"file": (name, text.encode("utf-8"), "text/plain")}).json()


def test_every_section_is_processed_with_source(client, env):
    upload(client, "uzun.txt", long_doc(4))

    def responder(messages, stream):
        if not stream:
            return summary()
        bolum, i, total, _ = bolum_of(messages)
        return jsonl({"tip": "kontrol", "kontrol_adi": f"Benzersiz kontrol {i} {bolum} zeta{i}", "kategori": "Operasyonel"})

    env.responder = responder
    s = wait_set(client, client.post("/api/sets").json()["id"])
    stream_calls = [c for c in env.calls if c.get("stream")]
    total = s["istatistik"]["bolum"]
    assert total >= 2 and len(stream_calls) == total
    assert s["istatistik"]["islenen"] == total and s["istatistik"]["basarisiz"] == 0
    assert len(s["controls"]) == total
    assert all(c["kaynak_belge"] == "uzun.txt" and c["kaynak_bolum"] for c in s["controls"])
    assert s["baslik"] == "Test Seti" and s["progress"] is None


def test_duplicates_across_sections_are_dropped(client, env):
    upload(client, "uzun.txt", long_doc(4))
    env.responder = lambda m, stream: jsonl(
        {"tip": "kontrol", "kontrol_adi": "Onay zincirinin bütünlüğünün kontrolü"},
        {"tip": "bosluk", "metin": "Hızlı onay yolu tanımlı değil."},
    ) if stream else summary()
    s = wait_set(client, client.post("/api/sets").json()["id"])
    assert len(s["controls"]) == 1
    assert s["istatistik"]["tekrar"] == s["istatistik"]["bolum"] - 1
    assert s["bosluklar"] == ["Hızlı onay yolu tanımlı değil."]


def test_quote_verification(client, env):
    upload(client, "p.txt", "1. Tahsis\n1.1. Her kredi başvurusu tahsis onayından önce risk analizinden geçer.\n")
    env.responder = lambda m, stream: jsonl(
        {"tip": "kontrol", "kontrol_adi": "Gerçek alıntılı", "kaynak": "tahsis onayından önce risk analizinden geçer"},
        {"tip": "kontrol", "kontrol_adi": "Uydurma alıntılı kontrol", "kaynak": "kredi komitesi her pazartesi tüm dosyaları onaylar"},
        {"tip": "kontrol", "kontrol_adi": "Alıntısız kontrol örneği", "kaynak": ""},
    ) if stream else summary()
    s = wait_set(client, client.post("/api/sets").json()["id"])
    by = {c["kontrol_adi"]: c["kaynak_dogrulandi"] for c in s["controls"]}
    assert by == {"Gerçek alıntılı": 1, "Uydurma alıntılı kontrol": 0, "Alıntısız kontrol örneği": None}
    assert s["istatistik"]["dogrulanamayan"] == 1


def test_one_bad_section_does_not_stop_the_rest(client, env):
    upload(client, "uzun.txt", long_doc(4))

    def responder(messages, stream):
        if not stream:
            return summary()
        _, i, _, _ = bolum_of(messages)
        if i == 1:
            return "Üzgünüm, bunu yapamam."              # biçim hatası
        if i == 2:
            return jsonl({"tip": "yok"})                  # meşru boş bölüm
        return jsonl({"tip": "kontrol", "kontrol_adi": f"Kontrol numara {i} omega{i}"})

    env.responder = responder
    s = wait_set(client, client.post("/api/sets").json()["id"])
    st = s["istatistik"]
    assert s["durum"] == "hazir"
    assert st["basarisiz"] == 1 and st["islenen"] == st["bolum"]
    assert len(s["controls"]) == st["bolum"] - 2
    assert s["hata"].startswith(f"1/{st['bolum']} bölüm işlenemedi")


def test_cancel_stops_generation(client, env):
    upload(client, "uzun.txt", long_doc(6))

    def slow(messages, stream):
        time.sleep(0.25)
        _, i, _, _ = bolum_of(messages) if stream else (0, 0, 0, "")
        return jsonl({"tip": "kontrol", "kontrol_adi": f"Yavaş kontrol {i} sigma{i}"}) if stream else summary()

    env.responder = slow
    set_id = client.post("/api/sets").json()["id"]
    time.sleep(0.3)
    assert client.post(f"/api/sets/{set_id}/cancel").json()["cancelled"]
    s = wait_set(client, set_id)
    assert s["durum"] == "iptal"
    assert s["istatistik"]["islenen"] < s["istatistik"]["bolum"]
    assert "durduruldu" in s["hata"]


def test_only_one_generation_at_a_time(client, env):
    upload(client, "uzun.txt", long_doc(3))
    env.responder = lambda m, stream: (time.sleep(0.2), jsonl({"tip": "yok"}))[1] if stream else summary()
    first = client.post("/api/sets").json()["id"]
    assert client.post("/api/sets").status_code == 409
    assert client.delete(f"/api/sets/{first}").status_code == 409
    client.post(f"/api/sets/{first}/cancel")
    wait_set(client, first)


def test_progress_returns_only_new_controls(client, env):
    upload(client, "p.txt", "1. A\nmetin\n")
    env.responder = lambda m, stream: jsonl({"tip": "kontrol", "kontrol_adi": "Tek kontrol"}) if stream else summary()
    s = wait_set(client, client.post("/api/sets").json()["id"])
    cid = s["controls"][0]["id"]
    assert len(client.get(f"/api/sets/{s['id']}/progress?after=0").json()["controls"]) == 1
    assert client.get(f"/api/sets/{s['id']}/progress?after={cid}").json()["controls"] == []


def test_unreadable_documents_are_skipped_and_reported(client, env):
    d = upload(client, "bos.txt", "   \n  ")
    assert d["uyari"] == "Belgede metin bulunamadı."
    upload(client, "iyi.txt", "1. A\nmetin burada\n")
    env.responder = lambda m, stream: jsonl({"tip": "kontrol", "kontrol_adi": "K"}) if stream else summary()
    s = wait_set(client, client.post("/api/sets").json()["id"])
    assert s["belge_adlari"] == ["iyi.txt"]


def test_recover_interrupted(env):
    with storage.db_session() as conn:
        conn.execute("INSERT INTO control_sets (id, baslik, durum) VALUES ('x1', 't', 'uretiliyor')")
    jobs.recover_interrupted()
    with storage.db_session() as conn:
        r = conn.execute("SELECT durum, hata FROM control_sets WHERE id='x1'").fetchone()
    assert r["durum"] == "hata" and "yeniden başlatıldı" in r["hata"]


def test_docx_headings_and_tables_extracted(tmp_path):
    doc = Document()
    doc.add_heading("Yetki Tablosu", level=1)
    doc.add_paragraph("Tahsis yetkileri aşağıdadır.")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Unvan", "Limit"
    t.cell(1, 0).text, t.cell(1, 1).text = "Şube Müdürü", "500.000 TL"
    p = tmp_path / "y.docx"
    doc.save(p)
    text, warn = documents.extract(p)
    assert warn is None
    assert "# Yetki Tablosu" in text and "Şube Müdürü | 500.000 TL" in text


def test_cp1254_text_file(tmp_path):
    p = tmp_path / "w.txt"
    p.write_bytes("Şube müdürü onaylar; İşlem kaydı tutulur.".encode("cp1254"))
    text, warn = documents.extract(p)
    assert warn is None and "Şube müdürü" in text and "İşlem" in text


def test_exports_include_source(client, env):
    upload(client, "p.txt", "1. Tahsis\n1.1. Her kredi başvurusu tahsis onayından önce risk analizinden geçer.\n")
    env.responder = lambda m, stream: jsonl(
        {"tip": "kontrol", "kontrol_adi": "Uydurma alıntılı", "kaynak": "kredi komitesi her pazartesi tüm dosyaları onaylar"},
    ) if stream else summary()
    s = wait_set(client, client.post("/api/sets").json()["id"])
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(client.get(f"/api/export/excel?set_id={s['id']}").content))["Kontroller"]
    header = [c.value for c in ws[1]]
    row = dict(zip(header, [c.value for c in ws[2]]))
    assert row["Kaynak belge"] == "p.txt" and row["Alıntı doğrulandı"].startswith("Hayır")
