"""Kurum kontrol kartı şablonu: KN.xx kodu, kart alanları, düzenleme ve Word/Excel çıktısı."""
import io
import json

from docx import Document
from openpyxl import load_workbook

from app import llm
from conftest import jsonl, wait_set

KART = {"tip": "kontrol", "kontrol_adi": "Sunuculara Erişimlerin Kontrolü", "kategori": "Bilgi Sistemleri",
        "hat": "2", "siklik": "Altı aylık", "risk": "Yüksek", "sorumlu": "İç Kontrol Müdürlüğü",
        "amac": "Domain Admin yetkilerinin altı ayda bir gözden geçirildiğini doğrulamak.",
        "nitelik": "İkinci seviye uygunluk ve işletim etkinliği kontrolü.",
        "ilgili_birimler": "Kimlik ve Erişim Yönetimi; sistem yöneticileri",
        "kapsam": "Kapsam içi tüm sunucular; tam ana kütle.",
        "donem": "Ocak–Haziran dönemi: Temmuz ayında; Temmuz–Aralık dönemi: izleyen yılın Ocak ayında",
        "adimlar": ["1) Gözden geçirme kayıtları incelenir.", "Gereksiz yetkilerin kaldırıldığı doğrulanır."],
        "kanit": "Sunucu envanteri; yetkili hesap listesi", "olcut": "Ana kütlenin tamamı gözden geçirilmiş olmalı.",
        "istisna": "-", "raporlama": "Altı aylık konsolide rapor."}


def _set(client, env):
    client.post("/api/documents", files={"file": ("e.txt", "1. Erişim\nYetkiler gözden geçirilir.\n".encode(), "text/plain")})
    env.responder = lambda m, stream: jsonl(KART, {**KART, "kontrol_adi": "İkinci kontrol", "adimlar": "a\nb"}) \
        if stream else json.dumps({"baslik": "Erişim", "ozet": "O"})
    return wait_set(client, client.post("/api/sets").json()["id"])


def test_prompt_asks_for_card_fields():
    for k in ("amac", "nitelik", "kapsam", "donem", "adimlar", "kanit", "olcut", "istisna", "raporlama"):
        assert f'"{k}"' in llm.SET_SYSTEM_PROMPT


def test_card_fields_stored_with_codes(client, env):
    s = _set(client, env)
    a, b = s["controls"]
    assert (a["kod"], b["kod"]) == ("KN.01", "KN.02")
    assert a["adimlar"] == "Gözden geçirme kayıtları incelenir.\nGereksiz yetkilerin kaldırıldığı doğrulanır."
    assert b["adimlar"] == "a\nb" and a["donem"].startswith("Ocak–Haziran") and a["kanit"]
    r = client.patch(f"/api/controls/{a['id']}", json={"kanit": "Yeni kanıt", "istisna": ""}).json()
    assert r["control"]["kanit"] == "Yeni kanıt" and r["control"]["istisna"] == ""


def test_word_uses_institution_card_template(client, env):
    s = _set(client, env)
    doc = Document(io.BytesIO(client.get(f"/api/export/word?set_id={s['id']}").content))
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "KN.01 – Sunuculara Erişimlerin Kontrolü" in heads and "KN.02 – İkinci kontrol" in heads
    t = doc.tables[0]
    labels = [r.cells[0].text for r in t.rows]
    assert labels == ["Kontrol Amacı", "Dayanak", "Kontrolün Niteliği", "İlgili Birimler", "Ana Kütle / Kapsam",
                      "Sıklık ve Dönem", "Uygulama Adımları", "Asgari Kanıt", "Değerlendirme Ölçütü",
                      "İstisna / Aksiyon", "Raporlama"]
    cell = {r.cells[0].text: r.cells[1].text for r in t.rows}
    assert cell["Sıklık ve Dönem"].startswith("6 aylık.\n• Ocak–Haziran dönemi: Temmuz ayında")
    assert cell["Uygulama Adımları"] == "1) Gözden geçirme kayıtları incelenir.\n2) Gereksiz yetkilerin kaldırıldığı doğrulanır."


def test_excel_and_library_have_card_columns(client, env):
    s = _set(client, env)
    ws = load_workbook(io.BytesIO(client.get(f"/api/export/excel?set_id={s['id']}").content))["Kontroller"]
    head = [c.value for c in ws[1]]
    assert head[:2] == ["Kod", "Kontrol"] and "Uygulama adımları" in head
    assert ws.cell(2, head.index("Uygulama adımları") + 1).value.startswith("1) Gözden")
    for c in s["controls"]:
        client.post(f"/api/controls/{c['id']}/decision", json={"karar": "kabul"})
    lib = client.get("/api/library").json()
    assert lib[0]["amac"] and lib[0]["kod"]
    doc = Document(io.BytesIO(client.get("/api/library/export/word").content))
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any(h.startswith("KN.01 – ") for h in heads) and any(h.startswith("KN.02 – ") for h in heads)
