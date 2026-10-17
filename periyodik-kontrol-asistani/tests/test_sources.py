import json

from conftest import jsonl, wait_set

from app import config, sources, storage

BS = """BANKALARIN BİLGİ SİSTEMLERİ YÖNETMELİĞİ
Amaç

MADDE 1 – (1) Bu Yönetmeliğin amacı bilgi sistemlerinin yönetimidir.
Kimlik ve erişim yönetimi

MADDE 11 – (1) Banka, kullanıcı erişim yetkilerini en az yılda bir gözden geçirir. Ayrıcalıklı hesapların kullanımı kayıt altına alınır ve izlenir.
İş sürekliliği

MADDE 13 – (1) Banka iş sürekliliği planını yılda bir test eder.
"""
AML = """TEDBİRLER YÖNETMELİĞİ
Şüpheli işlem

MADDE 27 – (1) Şüpheli işlemler MASAK'a bildirilir. Müşteri tanıma kuralları uygulanır.
Müşteri tanıma

MADDE 5 – (1) Yükümlüler müşterinin kimliğini tespit eder ve aklama riskini değerlendirir.
Kayıt

MADDE 46 – (1) Belgeler sekiz yıl saklanır.
"""


def _kaynaklar(tmp):
    d = tmp / "kaynaklar"
    (d / "metin").mkdir(parents=True)
    (d / "orijinal").mkdir()
    (d / "metin" / "bs.txt").write_text(BS, encoding="utf-8")
    (d / "metin" / "aml.txt").write_text(AML, encoding="utf-8")
    (d / "orijinal" / "bs.html").write_text("<html>BS</html>", encoding="utf-8")
    (d / "kaynaklar.json").write_text(json.dumps({
        "alanlar": ["Bilgi Sistemleri ve Veri Güvenliği", "AML / Uyum"],
        "kaynaklar": [
            {"id": "bs", "ad": "BS Yönetmeliği", "tur": "resmi", "kurum": "BDDK", "alanlar": ["Bilgi Sistemleri ve Veri Güvenliği"]},
            {"id": "aml", "ad": "Tedbirler Yönetmeliği", "tur": "resmi", "kurum": "MASAK", "alanlar": ["AML / Uyum"]},
            {"id": "yok", "ad": "Metni olmayan", "tur": "global", "alanlar": []},
        ]}, ensure_ascii=False), encoding="utf-8")
    (d / "indirme.json").write_text(json.dumps({"bs": {"url": "https://example", "erisim_tarihi": "2026-10-17",
                                                       "orijinal": "bs.html"}}), encoding="utf-8")
    with storage.db_session() as conn:
        sources.sync_builtin(conn)
        sources.sync_builtin(conn)    # ikinci eşitleme kopya oluşturmaz


def _builtin(client):
    return {d["filename"]: d for d in client.get("/api/documents").json() if d["yerlesik"]}


def test_builtin_sources_are_listed_and_protected(client, env, tmp_path):
    _kaynaklar(tmp_path)
    b = _builtin(client)
    assert set(b) == {"BS Yönetmeliği", "Tedbirler Yönetmeliği"}       # metni olmayan atlanır
    bs = b["BS Yönetmeliği"]
    assert bs["tur"] == "resmi" and bs["alanlar"] == ["Bilgi Sistemleri ve Veri Güvenliği"] and bs["bolum"] == 3
    assert client.delete(f"/api/documents/{bs['id']}").status_code == 400
    assert client.patch(f"/api/documents/{bs['id']}", json={"tur": "esas"}).status_code == 400
    t = client.get(f"/api/documents/{bs['id']}/text").json()
    assert [x["bolum"] for x in t["bolumler"]][1] == "Madde 11 — Kimlik ve erişim yönetimi"
    assert t["inline"] and client.get(f"/api/documents/{bs['id']}/file").text == "<html>BS</html>"
    # manifest'ten çıkan kaynak veritabanından da kalkar
    m = json.loads((tmp_path / "kaynaklar" / "kaynaklar.json").read_text())
    m["kaynaklar"] = m["kaynaklar"][:1]
    (tmp_path / "kaynaklar" / "kaynaklar.json").write_text(json.dumps(m))
    with storage.db_session() as conn:
        sources.sync_builtin(conn)
    assert set(_builtin(client)) == {"BS Yönetmeliği"}


def test_upload_suggests_type_and_area_and_user_confirms(client, env, tmp_path):
    _kaynaklar(tmp_path)
    esas = ("Erişim Uygulama Esasları.txt", "Kullanıcı erişim yetkileri yılda bir gözden geçirilir. "
            "Ayrıcalıklı hesapların kullanımı kayıt altına alınır ve haftalık izlenir.")
    d = client.post("/api/documents", files={"file": (esas[0], esas[1].encode(), "text/plain")}).json()
    assert d["tur"] == "esas" and d["alanlar"][0] == "Bilgi Sistemleri ve Veri Güvenliği" and d["onerildi"] == 1
    org = client.post("/api/documents", files={"file": ("Organizasyon ve Görev Tanımları.txt",
                      "Bilgi Güvenliği Müdürlüğü: erişim yetkilerini izler.".encode(), "text/plain")}).json()
    assert org["tur"] == "organizasyon"
    r = client.patch(f"/api/documents/{d['id']}", json={"alanlar": ["AML / Uyum", "Uydurma alan"]}).json()
    assert r["alanlar"] == ["AML / Uyum"] and r["onerildi"] == 0
    assert client.patch(f"/api/documents/{d['id']}", json={"tur": "resmi"}).status_code == 400


def test_scoped_generation_uses_basis_and_org(client, env, tmp_path):
    _kaynaklar(tmp_path)
    a = client.post("/api/documents", files={"file": ("erisim-esaslari.txt",
        "6. Ayrıcalıklı Hesaplar\nAyrıcalıklı hesap işlemleri Bilgi Güvenliği tarafından haftalık incelenir. "
        "Kullanıcı erişim yetkileri yılda bir gözden geçirilir.\n".encode(), "text/plain")}).json()
    other = client.post("/api/documents", files={"file": ("kredi-esaslari.txt",
        "1. Kredi\nKrediler kredi komitesince onaylanır.\n".encode(), "text/plain")}).json()
    client.post("/api/documents", files={"file": ("Organizasyon Şeması ve Görev Tanımları.txt",
        "Bilgi Güvenliği Müdürlüğü ayrıcalıklı hesap işlemlerini inceler.".encode(), "text/plain")})
    seen = {}

    def responder(messages, stream):
        if not stream:
            return json.dumps({"baslik": "Model başlığı", "ozet": "Özet"})
        seen["user"] = messages[-1]["content"]
        return jsonl({"tip": "kontrol", "kontrol_adi": "Ayrıcalıklı hesap incelemesinin teyidi", "hat": "2",
                      "dayanak_no": [1], "dayanak": "", "sorumlu": "Bilgi Güvenliği Müdürlüğü",
                      "kaynak": "Ayrıcalıklı hesap işlemleri Bilgi Güvenliği tarafından haftalık incelenir."})
    env.responder = responder
    bs = _builtin(client)["BS Yönetmeliği"]
    body = {"belgeler": [a["id"]], "dayanaklar": [bs["id"]], "organizasyon": True,
            "alanlar": ["Bilgi Sistemleri ve Veri Güvenliği"], "baslik": "Veri Güvenliği — Ekim 2026"}
    s = wait_set(client, client.post("/api/sets", json=body).json()["id"])
    assert s["baslik"] == "Veri Güvenliği — Ekim 2026" and s["ozet"] == "Özet"     # kullanıcının adı korunur
    assert s["belge_adlari"] == ["erisim-esaslari.txt"]                               # kredi işlenmedi
    assert s["kapsam"]["dayanak"] == [bs["id"]] and s["kapsam"]["organizasyon"]
    assert "DAYANAK ADAYLARI" in seen["user"] and "Madde 11" in seen["user"]
    assert "ŞİRKET ORGANİZASYONU" in seen["user"] and "Bilgi Güvenliği Müdürlüğü" in seen["user"]
    c = s["controls"][0]
    assert c["sorumlu"] == "Bilgi Güvenliği Müdürlüğü"
    assert c["dayanak_bag"][0]["doc_id"] == bs["id"] and c["dayanak_bag"][0]["label"].startswith("Madde 11")
    assert "BS Yönetmeliği — Madde 11" in c["dayanak"]          # model boş bıraktı: bağlantıdan dolduruldu
    assert _builtin(client)["BS Yönetmeliği"]["kullanim"] == 1
    assert {d["filename"]: d for d in client.get("/api/documents").json()}["kredi-esaslari.txt"]["kullanim"] == 0
    # geçersiz seçimler
    assert client.post("/api/sets", json={"belgeler": [bs["id"]]}).status_code == 400
    assert client.post("/api/sets", json={"belgeler": [other["id"]], "dayanaklar": [a["id"]]}).status_code == 400
    assert client.post("/api/sets", json={"belgeler": []}).status_code == 400


def test_dayanak_by_madde_number_without_dayanak_no(client, env, tmp_path):
    _kaynaklar(tmp_path)
    a = client.post("/api/documents", files={"file": ("e.txt",
        "1. Erişim\nErişim yetkileri yılda bir gözden geçirilir; ayrıcalıklı hesaplar izlenir.\n".encode(), "text/plain")}).json()
    env.responder = lambda m, stream: jsonl({"tip": "kontrol", "kontrol_adi": "Yetki gözden geçirme",
                                             "dayanak": "BDDK BS Yönetmeliği m.11"}) if stream else "{}"
    bs = _builtin(client)["BS Yönetmeliği"]
    s = wait_set(client, client.post("/api/sets", json={"belgeler": [a["id"]], "dayanaklar": [bs["id"]]}).json()["id"])
    c = s["controls"][0]
    assert c["dayanak"] == "BDDK BS Yönetmeliği m.11" and c["dayanak_bag"][0]["label"].startswith("Madde 11")


def test_sorumlu_editable_and_exported(client, env):
    env.responder = lambda m, stream: jsonl({"tip": "kontrol", "kontrol_adi": "K", "sorumlu": "Şube"}) if stream else "{}"
    client.post("/api/documents", files={"file": ("a.txt", b"1. A\nSube islemi onaylar.\n", "text/plain")})
    s = wait_set(client, client.post("/api/sets").json()["id"])
    cid = s["controls"][0]["id"]
    r = client.patch(f"/api/controls/{cid}", json={"sorumlu": "Operasyon Müdürlüğü"})
    assert r.status_code == 200
    client.post(f"/api/controls/{cid}/decision", json={"karar": "kabul"})
    x = client.get(f"/api/export/excel?set_id={s['id']}")
    from io import BytesIO
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(x.content)).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    head = next(r for r in rows if "Sorumlu" in r)
    assert rows[rows.index(head) + 1][head.index("Sorumlu")] == "Operasyon Müdürlüğü"


def test_meta_has_areas_and_types(client, env):
    m = client.get("/api/meta").json()
    assert m["turler"]["esas"] == "Uygulama esası" and m["alanlar"]
