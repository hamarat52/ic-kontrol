"""Arayüz v2 için uç noktalar: toplu karar, ret geri alma, belge bilgileri, öğrenilenleri yönetme."""
import json

from app import config, skills
from conftest import jsonl, wait_set


def _set(client, env, n=3):
    client.post("/api/documents", files={"file": ("p.txt", "1. A\nmetin burada\n".encode(), "text/plain")})
    names = ["Teminat değerleme raporu", "Kredi limit aşımı izlemesi", "Mutabakat farkı bildirimi"][:n]
    env.responder = lambda m, stream: jsonl(*[{"tip": "kontrol", "kontrol_adi": x, "kategori": "Operasyonel"} for x in names]) \
        if stream else json.dumps({"baslik": "B", "ozet": "O"})
    return wait_set(client, client.post("/api/sets").json()["id"])


def _learned(name="kontrol-metrikleri"):
    return (config.SKILLS_DIR / f"{name}.md").read_text(encoding="utf-8")


def test_bulk_decision(client, env):
    s = _set(client, env)
    ids = [c["id"] for c in s["controls"]]
    r = client.post("/api/controls/bulk-decision", json={"ids": ids[:2], "karar": "kabul"}).json()
    assert [c["status"] for c in r["controls"]] == ["kabul", "kabul"]
    assert len(client.get("/api/library").json()) == 2
    assert client.post("/api/controls/bulk-decision", json={"ids": [], "karar": "kabul"}).status_code == 400
    assert client.post("/api/controls/bulk-decision", json={"ids": ids, "karar": "x"}).status_code == 400
    assert client.post("/api/controls/bulk-decision", json={"ids": [999999], "karar": "red"}).status_code == 404


def test_undoing_reject_removes_its_learning(client, env):
    s = _set(client, env, 1)
    cid = s["controls"][0]["id"]
    client.post(f"/api/controls/{cid}/decision", json={"karar": "red"})
    assert "Reddedildi: 'Teminat değerleme raporu'" in _learned()
    client.post(f"/api/controls/{cid}/decision", json={"karar": "bekliyor"})
    assert "Reddedildi: 'Teminat değerleme raporu'" not in _learned()
    arsiv = (config.DATA_DIR / "skills_arsiv.jsonl").read_text(encoding="utf-8")
    assert "Teminat değerleme raporu" in arsiv


def test_edit_and_archive_learning(client, env):
    skills.append_learning("kontrol-metrikleri", "Şube hacmi düşük; sıklıklar haftalıktan sık olmasın.")
    r = client.patch("/api/skills/kontrol-metrikleri/learnings",
                     json={"eski": "Şube hacmi düşük; sıklıklar haftalıktan sık olmasın.", "yeni": "Kritik riskler hariç sıklık en fazla haftalık."})
    assert r.status_code == 200
    md = _learned()
    assert "Kritik riskler hariç sıklık en fazla haftalık." in md and "Şube hacmi düşük" not in md
    assert client.post("/api/skills/kontrol-metrikleri/learnings/archive",
                       json={"eski": "Kritik riskler hariç sıklık en fazla haftalık."}).status_code == 200
    assert "Kritik riskler hariç" not in _learned()
    # güvenlik ve doğrulama
    assert client.patch("/api/skills/..%2Fapp/learnings", json={"eski": "x", "yeni": "y"}).status_code == 404
    assert client.patch("/api/skills/yok-boyle/learnings", json={"eski": "x", "yeni": "y"}).status_code == 404
    assert client.patch("/api/skills/kontrol-metrikleri/learnings", json={"eski": "olmayan", "yeni": "y"}).status_code == 400
    assert client.patch("/api/skills/kontrol-metrikleri/learnings", json={"eski": "x", "yeni": "  "}).status_code == 400


def test_documents_metadata(client, env):
    body = "1. Amaç\nKredi süreci.\n2. Onay\nŞube müdürü onaylar.\n".encode()
    a = client.post("/api/documents", files={"file": ("a.txt", body, "text/plain")}).json()
    assert a["bolum"] == 1 and a["kopya_of"] is None and a["kullanim"] == 0
    b = client.post("/api/documents", files={"file": ("a (1).txt", body, "text/plain")}).json()
    assert b["kopya_of"] == "a.txt"
    env.responder = lambda m, stream: jsonl({"tip": "kontrol", "kontrol_adi": "K"}) if stream else json.dumps({"baslik": "B", "ozet": "O"})
    wait_set(client, client.post("/api/sets").json()["id"])
    docs = {d["filename"]: d for d in client.get("/api/documents").json()}
    assert docs["a.txt"]["kullanim"] == 1 and docs["a (1).txt"]["kullanim"] == 0   # kopya üretimde atlanır
    assert len([c for c in env.calls if c.get("stream")]) == 1


def test_document_view_and_download(client, env):
    body = "1. Amaç\nKredi süreci.\n2. Onay\nŞube müdürü onaylar.\n".encode()
    d = client.post("/api/documents", files={"file": ("Kredi Prosedürü.txt", body, "text/plain")}).json()
    r = client.get(f"/api/documents/{d['id']}/text").json()
    assert r["filename"] == "Kredi Prosedürü.txt" and r["inline"] is True
    assert "Şube müdürü onaylar." in "".join(b["text"] for b in r["bolumler"])
    f = client.get(f"/api/documents/{d['id']}/file?indir=1")
    assert f.status_code == 200 and f.content == body
    assert f.headers["content-disposition"].startswith("attachment;") and "Kredi%20Prosed" in f.headers["content-disposition"]
    assert client.get(f"/api/documents/{d['id']}/file").headers["content-disposition"].startswith("inline;")
    assert client.get("/api/documents/99999/text").status_code == 404
    assert client.get("/api/documents/99999/file").status_code == 404


def test_library_export(client, env):
    import io
    from docx import Document
    from openpyxl import load_workbook
    assert client.get("/api/library/export/word").status_code == 400
    s = _set(client, env)
    ids = [c["id"] for c in s["controls"]]
    client.post("/api/controls/bulk-decision", json={"ids": ids[:2], "karar": "kabul"})
    w = client.get("/api/library/export/word")
    assert w.status_code == 200 and "Kontrol%20K%C3%BCt%C3%BCphanesi" in w.headers["content-disposition"]
    text = "\n".join(p.text for p in Document(io.BytesIO(w.content)).paragraphs)
    assert "Kontrol Kütüphanesi" in text and "2 kontrol" in text
    ws = load_workbook(io.BytesIO(client.get("/api/library/export/excel").content))["Kontroller"]
    header = [c.value for c in ws[1]]
    assert header[-1] == "Kaynak set" and "Durum" not in header
    assert ws.max_row == 3
    row = dict(zip(header, [c.value for c in ws[2]]))
    assert row["Kaynak set"] == "B" and row["Kaynak belge"] == "p.txt"
